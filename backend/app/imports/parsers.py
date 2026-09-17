"""Turn a bank's CSV/Excel export into normalized rows.

Adding a bank means adding a `ColumnMap` preset, not writing a parser. Every
preset normalizes to Plaid's sign convention (**positive = money out**) so
imported rows are indistinguishable from synced ones downstream.
"""
from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime


class ParseError(ValueError):
    """The file could not be read as a statement under the chosen preset."""


@dataclass(frozen=True)
class ParsedRow:
    date: date
    description: str
    amount: float  # normalized: positive = money out
    currency: str | None


@dataclass(frozen=True)
class ColumnMap:
    """How one bank's columns map onto `ParsedRow`.

    Header matching is case/whitespace-insensitive. `date` and `amount` are
    *alternatives* (first present, non-empty column wins). `description` is a
    *join* — every listed column that's present is concatenated, which is how
    banks that split a memo across "Description 1"/"Description 2" read back.
    """

    label: str
    date: tuple[str, ...]
    description: tuple[str, ...]
    date_formats: tuple[str, ...]
    amount: tuple[str, ...] = ()
    # Separate debit/credit columns, used when `amount` is empty (common in
    # Excel statement exports). Debit = money out.
    debit: tuple[str, ...] = ()
    credit: tuple[str, ...] = ()
    # True when the file writes money-out as a negative number, i.e. the
    # opposite of Plaid, so the sign is flipped on read.
    debit_negative: bool = True
    # Currency implied by which amount column matched, e.g. {"cad$": "CAD"}.
    currency_by_column: dict[str, str] = field(default_factory=dict)
    default_currency: str | None = None


RBC = ColumnMap(
    label="RBC (personal banking export)",
    date=("Transaction Date", "Date"),
    description=("Description 1", "Description 2"),
    amount=("CAD$", "USD$"),
    date_formats=("%m/%d/%Y", "%Y-%m-%d", "%d-%b-%Y"),
    debit_negative=True,
    currency_by_column={"cad$": "CAD", "usd$": "USD"},
    default_currency="CAD",
)

GENERIC = ColumnMap(
    label="Generic (common column names)",
    date=("Date", "Transaction Date", "Posted Date", "Posting Date"),
    description=("Description", "Details", "Memo", "Narrative", "Payee"),
    amount=("Amount", "Value"),
    debit=("Debit", "Withdrawal", "Withdrawals"),
    credit=("Credit", "Deposit", "Deposits"),
    date_formats=("%Y-%m-%d", "%m/%d/%Y", "%d/%m/%Y", "%d-%b-%Y", "%b %d, %Y"),
    debit_negative=True,
)

PRESETS: dict[str, ColumnMap] = {"rbc": RBC, "generic": GENERIC}


def _norm(header: str) -> str:
    return re.sub(r"\s+", " ", (header or "").replace("﻿", "").strip()).lower()


def _pick(row: dict[str, str], candidates: tuple[str, ...]) -> tuple[str, str] | None:
    """First candidate column present with a non-empty value, as (column, value)."""
    for candidate in candidates:
        key = _norm(candidate)
        if key in row and str(row[key]).strip():
            return key, str(row[key]).strip()
    return None


_MONEY_JUNK = re.compile(r"[^\d.\-()]")


def _parse_amount(raw: str) -> float:
    """Money text to float. Handles $, thousands separators and (123.45) negatives."""
    text = _MONEY_JUNK.sub("", raw)
    negative = text.startswith("(") and text.endswith(")")
    text = text.strip("()")
    if not text or text in {"-", "."}:
        raise ParseError(f"Unreadable amount: {raw!r}")
    try:
        value = float(text)
    except ValueError as exc:
        raise ParseError(f"Unreadable amount: {raw!r}") from exc
    return -value if negative else value


def _parse_date(raw: object, formats: tuple[str, ...]) -> date:
    if isinstance(raw, datetime):
        return raw.date()
    if isinstance(raw, date):
        return raw
    text = str(raw).strip()
    for fmt in formats:
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    try:
        return datetime.fromisoformat(text).date()
    except ValueError as exc:
        raise ParseError(f"Unreadable date: {raw!r}") from exc


def _row_amount(row: dict[str, str], cmap: ColumnMap) -> tuple[float, str | None]:
    """Signed amount in Plaid's convention, plus the currency the column implies."""
    hit = _pick(row, cmap.amount) if cmap.amount else None
    if hit is not None:
        column, raw = hit
        amount = _parse_amount(raw)
        if cmap.debit_negative:
            amount = -amount
        currency = cmap.currency_by_column.get(column, cmap.default_currency)
        return amount, currency

    debit = _pick(row, cmap.debit) if cmap.debit else None
    if debit is not None:
        return abs(_parse_amount(debit[1])), cmap.default_currency
    credit = _pick(row, cmap.credit) if cmap.credit else None
    if credit is not None:
        return -abs(_parse_amount(credit[1])), cmap.default_currency

    raise ParseError("Row has no amount")


def _row_description(row: dict[str, str], cmap: ColumnMap) -> str:
    parts = [
        str(row[key]).strip()
        for key in (_norm(c) for c in cmap.description)
        if key in row and str(row[key]).strip()
    ]
    return " ".join(parts)


def _read_csv(data: bytes) -> list[list[str]]:
    text = data.decode("utf-8-sig", errors="replace")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    return [row for row in csv.reader(io.StringIO(text), dialect) if any(c.strip() for c in row)]


def _read_xlsx(data: bytes) -> list[list[object]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - dependency is declared
        raise ParseError("Excel support requires the openpyxl package") from exc

    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sheet = workbook.active
    rows = [
        list(row)
        for row in sheet.iter_rows(values_only=True)
        if any(c is not None and str(c).strip() for c in row)
    ]
    workbook.close()
    return rows


def _find_header(rows: list[list], cmap: ColumnMap) -> int:
    """Index of the header row — banks often prepend account/preamble lines."""
    wanted = {_norm(c) for c in cmap.date}
    for index, row in enumerate(rows[:20]):
        if wanted & {_norm(str(cell)) for cell in row if cell is not None}:
            return index
    raise ParseError(
        f"No header row found. Expected a date column named one of: {', '.join(cmap.date)}"
    )


def parse(data: bytes, filename: str, cmap: ColumnMap) -> list[ParsedRow]:
    """Parse an uploaded statement. Raises ParseError if the shape doesn't fit."""
    rows = _read_xlsx(data) if filename.lower().endswith((".xlsx", ".xlsm")) else _read_csv(data)
    if not rows:
        raise ParseError("File is empty")

    header_index = _find_header(rows, cmap)
    headers = [_norm(str(cell)) if cell is not None else "" for cell in rows[header_index]]

    parsed: list[ParsedRow] = []
    for line_number, raw_row in enumerate(rows[header_index + 1 :], start=header_index + 2):
        record = {h: raw_row[i] for i, h in enumerate(headers) if h and i < len(raw_row)}
        record = {k: ("" if v is None else v) for k, v in record.items()}
        date_hit = _pick({k: str(v) for k, v in record.items()}, cmap.date)
        if date_hit is None:
            continue  # trailing totals/blank rows
        try:
            when = _parse_date(record[date_hit[0]], cmap.date_formats)
            amount, currency = _row_amount({k: str(v) for k, v in record.items()}, cmap)
        except ParseError as exc:
            raise ParseError(f"Line {line_number}: {exc}") from exc
        parsed.append(
            ParsedRow(
                date=when,
                description=_row_description({k: str(v) for k, v in record.items()}, cmap),
                amount=amount,
                currency=currency,
            )
        )

    if not parsed:
        raise ParseError("No transaction rows found")
    return parsed
